#!/usr/bin/env python3
"""Build deterministic workshop sample files for the messy archive exercise."""

from __future__ import annotations

import csv
import math
import random
import shutil
import wave
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


REPO_ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = REPO_ROOT / "sample-files" / "messy-research-archive"


def reset_archive() -> None:
    if ARCHIVE.exists():
        shutil.rmtree(ARCHIVE)
    ARCHIVE.mkdir(parents=True)


def write_text(path: Path, body: str) -> None:
    path.write_text(body.strip() + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    wb.save(path)


def write_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def write_pdf(path: Path, title: str, paragraphs: list[str]) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, title)
    c.setFont("Helvetica", 10)
    y = height - 110
    for paragraph in paragraphs:
        for line in wrap(paragraph, 88):
            c.drawString(72, y, line)
            y -= 14
            if y < 80:
                c.showPage()
                c.setFont("Helvetica", 10)
                y = height - 72
        y -= 10
    c.save()


def wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line: list[str] = []
    for word in words:
        if sum(len(part) for part in line) + len(line) + len(word) > width:
            lines.append(" ".join(line))
            line = [word]
        else:
            line.append(word)
    if line:
        lines.append(" ".join(line))
    return lines


def write_image(path: Path, size: tuple[int, int], seed: int, label: str) -> None:
    rng = random.Random(seed)
    image = Image.new("RGB", size, (242, 242, 236))
    draw = ImageDraw.Draw(image)
    for _ in range(420):
        x0 = rng.randrange(0, size[0])
        y0 = rng.randrange(0, size[1])
        x1 = min(size[0], x0 + rng.randrange(20, 220))
        y1 = min(size[1], y0 + rng.randrange(10, 160))
        colour = (
            rng.randrange(35, 215),
            rng.randrange(45, 205),
            rng.randrange(50, 220),
        )
        draw.rectangle((x0, y0, x1, y1), outline=colour, width=rng.randrange(1, 4))
    draw.rectangle((24, 24, size[0] - 24, 86), fill=(250, 250, 246), outline=(40, 40, 40))
    draw.text((44, 46), label, fill=(20, 20, 20))
    image.save(path, quality=88)


def write_wav(path: Path) -> None:
    sample_rate = 16000
    duration_seconds = 45
    with wave.open(str(path), "w") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(2)
        handle.setframerate(sample_rate)
        for i in range(sample_rate * duration_seconds):
            value = int(12000 * math.sin(2 * math.pi * 440 * i / sample_rate))
            handle.writeframesraw(value.to_bytes(2, byteorder="little", signed=True))


def write_noise_png(path: Path, size: tuple[int, int], seed: int) -> None:
    rng = random.Random(seed)
    pixels = bytes(rng.randrange(0, 256) for _ in range(size[0] * size[1]))
    image = Image.frombytes("L", size, pixels)
    image.save(path)


def write_large_mock_csv(path: Path) -> None:
    rng = random.Random(20260610)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["record_id", "timestamp", "sensor_a", "sensor_b", "status", "comment"])
        for i in range(60000):
            writer.writerow(
                [
                    f"M{i:06}",
                    f"2025-03-{1 + (i % 28):02}T{(i // 3600) % 24:02}:{(i // 60) % 60:02}:{i % 60:02}",
                    round(rng.uniform(10, 90), 4),
                    round(rng.uniform(0, 1), 6),
                    ["ok", "check", "repeat", "missing"][i % 4],
                    "fictional high-volume mock export",
                ]
            )


def build() -> None:
    reset_archive()
    write_text(
        ARCHIVE / "README.md",
        """
        # Messy Research Archive

        This is a fictional mock research archive. It is intentionally untidy.

        All files are in this one folder. There are no helpful subfolders because
        the participant task is to ask Codex to inspect the material, infer useful
        categories, and propose a safer organisation before moving anything.

        ## What Is Inside

        - CSV files with similar names, old dates, and overlapping rows.
        - Excel workbooks with old or unclear labels.
        - Word documents that mix protocols, notes, and draft findings.
        - Text notes and transcript fragments.
        - PDF placeholders.
        - PNG and JPG image files.
        - A small WAV placeholder.
        - Larger mock exports that make the folder feel more like a real data archive.

        All content is fictional. The files imitate common research-project messiness
        without using real participant data or sensitive material.

        ## Suggested Exercise

        Ask Codex to:

        1. Build a file inventory with names, formats, sizes, and likely purposes.
        2. Flag duplicate-looking files and unclear file names.
        3. Propose a safer folder structure.
        4. Draft a `data-inventory.md`.
        5. Explain what it would move or rename before making any changes.

        Participants should review the plan before asking Codex to change files.
        """,
    )

    survey_rows = [
        {
            "response_id": f"R{i:03}",
            "wave": "pilot" if i < 8 else "main",
            "group": ["control", "intervention", "unknown"][i % 3],
            "confidence_pre": 2 + (i % 4),
            "confidence_post": 3 + (i % 5),
            "free_text": [
                "Useful but I need clearer steps.",
                "The folder was hard to understand.",
                "I liked having outputs to inspect.",
                "Not sure which file is final.",
            ][i % 4],
        }
        for i in range(1, 31)
    ]
    write_csv(ARCHIVE / "survey_export_2025-03-04.csv", survey_rows)
    write_csv(ARCHIVE / "survey_export_FINAL_final.csv", list(reversed(survey_rows[:24])))
    write_csv(ARCHIVE / "2024-11-18_survey_export_old-platform.csv", survey_rows[:18])
    write_csv(ARCHIVE / "copy of survey export 2025.csv", survey_rows[6:30])
    write_csv(
        ARCHIVE / "participant_lookup_DO_NOT_USE_mock.csv",
        [
            {"participant_code": f"P{i:02}", "consent": "yes", "notes": "fictional mock row"}
            for i in range(1, 13)
        ],
    )

    write_workbook(
        ARCHIVE / "coding_export_nodes_v3.xlsx",
        {
            "Nodes": [
                ["node", "count", "review_status"],
                ["setup friction", 14, "needs merge"],
                ["documentation", 22, "reviewed"],
                ["publishing uncertainty", 9, "unclear"],
            ],
            "Memos": [
                ["memo_id", "summary"],
                ["M01", "Several notes refer to the same setup issue."],
                ["M02", "Final labels may not match transcript headings."],
            ],
        },
    )
    write_workbook(
        ARCHIVE / "analysis_workbook_old_DO_NOT_USE.xlsx",
        {
            "Sheet1": [
                ["metric", "value"],
                ["mean_pre", 3.2],
                ["mean_post", 4.1],
                ["missing_rows", 6],
            ]
        },
    )
    write_workbook(
        ARCHIVE / "2023 pilot coding matrix FINAL.xlsx",
        {
            "codes": [
                ["case", "code", "certainty"],
                ["A01", "setup friction", "medium"],
                ["A02", "documentation", "high"],
                ["A03", "unclear output ownership", "low"],
            ],
            "notes": [["created", "2023-12-08"], ["status", "superseded?"]],
        },
    )

    write_text(
        ARCHIVE / "Interview notes clean FINAL v2.txt",
        """
        Interview A notes. The participant described using a local folder, then losing
        track of which transcript had been cleaned. They wanted the agent to explain
        its file changes before making them. This is fictional mock content.
        """,
    )
    write_text(
        ARCHIVE / "transcript_A_raw_unchecked.txt",
        """
        SPEAKER 1: I put everything in Downloads first.
        SPEAKER 2: Then what happened?
        SPEAKER 1: I made a folder called final but it has old files too.
        [inaudible]
        """,
    )
    write_text(
        ARCHIVE / "README maybe for archive.txt",
        """
        Notes from an imaginary assistant handover. Some file names are misleading on
        purpose. No file contains real participant data.
        """,
    )
    write_text(
        ARCHIVE / "draft-findings-notes.md",
        """
        # Draft findings

        These notes are deliberately rough. The exercise is to ask an agent to inspect
        the folder, build an inventory, identify duplicates, and propose a cleaner
        structure before moving or renaming anything.
        """,
    )
    write_text(
        ARCHIVE / "2023-10-02 fieldnotes maybe duplicate.txt",
        """
        Short fictional fieldnote. Mentions that the first pilot run used paper forms,
        but does not say whether the forms were later transcribed.
        """,
    )
    write_text(
        ARCHIVE / "notes_from_meeting_14 Jan 2024.txt",
        """
        Project meeting notes. Action: check whether the 2024 export replaced the 2023
        pilot coding matrix. Action owner not recorded.
        """,
    )
    write_text(
        ARCHIVE / "transcript B edited maybe.txt",
        """
        SPEAKER 1: The agent should ask before it moves files.
        SPEAKER 2: Especially when file names say final.
        SPEAKER 1: Yes, final can just mean I was tired.
        """,
    )
    write_text(
        ARCHIVE / "analysis notes 2025-02-28.md",
        """
        # Analysis notes

        Compare pre and post confidence scores. Check whether pilot rows are included
        twice. Do not use the participant lookup file in shared outputs.
        """,
    )
    write_text(
        ARCHIVE / "tmp-summary-latest.md",
        """
        # Temporary summary

        This summary may have been generated from an old export. It should be checked
        against the current CSV files before reuse.
        """,
    )
    write_text(
        ARCHIVE / "2025-03-01 TODO after workshop.txt",
        """
        - Rename confusing files.
        - Check duplicate survey exports.
        - Separate raw data, derived outputs, and admin files.
        """,
    )
    write_text(
        ARCHIVE / "untitled notes from desktop.txt",
        """
        Not sure whether these notes belong with interview A or B. The source folder was
        probably Downloads.
        """,
    )
    write_text(
        ARCHIVE / "2022-12-15 archive-index-from-old-laptop.csv",
        """
        filename,maybe_description,status
        pilot_notes.docx,old pilot notes,unknown
        data_export.csv,raw export maybe,superseded
        figure-final.png,chart for report,unknown
        """,
    )
    write_text(
        ARCHIVE / "table_export_2023-12-01.tsv",
        """
        case_id\twave\tcode\tcomment
        P01\tpilot\tsetup friction\tfictional row
        P02\tpilot\tdocumentation\tfictional row
        P03\tpilot\tpublishing uncertainty\tfictional row
        """,
    )
    write_text(
        ARCHIVE / "duplicate transcript A 2024 maybe older.txt",
        """
        SPEAKER 1: I put everything into downloads first.
        SPEAKER 2: Then what happened?
        SPEAKER 1: I created a final folder but it also has old files.
        """,
    )
    write_text(
        ARCHIVE / "chat export with RA 2024-02-19.txt",
        """
        10:14 RA: I found another survey export.
        10:16 PI: Please do not delete anything yet.
        10:22 RA: Should I put it with the old files?
        """,
    )
    write_text(
        ARCHIVE / "analysis-output-copy-pasted-from-email.md",
        """
        # Copied output

        This appears to be a result summary pasted from an email. It has no clear link
        to a script or input file.
        """,
    )

    write_pdf(
        ARCHIVE / "ethics-draft-final-v2.pdf",
        "Mock ethics draft",
        [
            "This fictional PDF imitates a short administrative document. It exists so the archive contains a realistic mix of file formats.",
            "The file should not be treated as a completed ethics document. It is placeholder content for a folder-organisation exercise.",
        ],
    )
    write_pdf(
        ARCHIVE / "scan-from-printer-001.pdf",
        "Mock scanned note",
        [
            "A short scanned-note placeholder. The vague name is intentional because participants will practise asking Codex to infer likely categories from content.",
        ],
    )
    write_pdf(
        ARCHIVE / "2023-09-30 consent form old version.pdf",
        "Mock consent form old version",
        [
            "This fictional PDF represents an old administrative form. It should be classified separately from research data.",
        ],
    )
    write_pdf(
        ARCHIVE / "exported report 2025-03-07 no appendix.pdf",
        "Mock exported report",
        [
            "This fictional report appears to summarise analysis results, but it does not state which input files were used.",
        ],
    )

    write_docx(
        ARCHIVE / "interview_protocol_2024-01-12.docx",
        "Mock interview protocol",
        [
            "This fictional protocol lists draft questions for a pilot interview.",
            "The document is included so participants must distinguish instruments from data and outputs.",
        ],
    )
    write_docx(
        ARCHIVE / "Findings draft FINAL to share.docx",
        "Draft findings",
        [
            "This fictional document looks polished but should be checked against the source data before being treated as final.",
            "The title deliberately uses an unreliable status label.",
        ],
    )
    write_docx(
        ARCHIVE / "2025-03-06 meeting notes with supervisor.docx",
        "Supervisor meeting notes",
        [
            "Check whether the old analysis workbook has been superseded.",
            "Create a clear data inventory before moving files.",
        ],
    )
    write_docx(
        ARCHIVE / "old lit review notes 2023.docx",
        "Old literature notes",
        [
            "These fictional notes may belong in documentation rather than the analysis data folder.",
        ],
    )
    write_docx(
        ARCHIVE / "2024_03_15 coding meeting messy notes.docx",
        "Coding meeting messy notes",
        [
            "The team discussed merging setup friction with onboarding difficulty.",
            "No final coding decision is recorded in this document.",
        ],
    )
    write_docx(
        ARCHIVE / "QUESTIONNAIRE old version maybe used.docx",
        "Questionnaire old version",
        [
            "This fictional questionnaire may or may not match the survey export.",
            "Participants should ask Codex to identify uncertainty rather than assume it is current.",
        ],
    )

    write_image(ARCHIVE / "Figure 1 final FINAL.png", (1600, 1000), 11, "Mock figure final FINAL")
    write_image(ARCHIVE / "IMG_0042.JPG", (2400, 1600), 42, "Mock field image")
    write_image(ARCHIVE / "screenshot-data-cleaning.png", (1400, 900), 87, "Mock screenshot")
    write_image(ARCHIVE / "2024-02-11 whiteboard photo.jpg", (1800, 1200), 64, "Mock whiteboard")
    write_image(ARCHIVE / "plot_new_new2.png", (1200, 800), 73, "Mock plot new new2")
    write_image(ARCHIVE / "figure2_draft_2024-12-08.png", (1100, 700), 91, "Mock figure 2 draft")
    write_wav(ARCHIVE / "meeting_audio_clip_placeholder.wav")
    write_noise_png(ARCHIVE / "raw-camera-export-not-reviewed.png", (2600, 1800), 19)
    write_noise_png(ARCHIVE / "scan_batch_07_uncropped.png", (2200, 1600), 31)
    write_large_mock_csv(ARCHIVE / "sensor_export_all_rows_maybe_final.csv")


if __name__ == "__main__":
    build()
